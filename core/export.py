"""XLSX export for jobs and all sellers."""
import io

from django import http
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.models import CollectionJob, Seller

HEADERS = [
    "Маркетплейс", "Страница продавца", "Название продавца", "Категории",
    "Мобильные телефоны", "Telegram", "Городские телефоны", "Email", "Сайт",
    "Срок работы на маркетплейсе (лет)", "Средняя оценка", "Дата создания",
    "Юридический адрес", "ИНН", "Город", "Дата обновления",
]


def _seller_row(s: Seller) -> list:
    return [
        s.get_marketplace_display(),
        s.seller_url or "",
        s.name or "",
        "; ".join(s.category_titles),
        ", ".join(s.mobile_phones),
        s.telegram_link or "",
        ", ".join(s.city_phones),
        ", ".join(s.emails),
        s.website or "",
        s.years_on_marketplace if s.years_on_marketplace is not None else "",
        s.rating or "",
        s.registered_at.strftime("%d.%m.%Y") if s.registered_at else "",
        s.legal_address or "",
        s.inn or "",
        s.city.name if s.city_id else "",
        timezone_str(s.last_updated_at),
    ]


def timezone_str(dt):
    from django.utils import timezone as tz

    return tz.localtime(dt).strftime("%d.%m.%Y %H:%M") if dt else ""


def build_xlsx(sellers) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Продавцы"
    ws.append(HEADERS)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    for s in sellers:
        ws.append(_seller_row(s))
    for idx, width in enumerate([14, 40, 32, 28, 26, 22, 26, 28, 24, 12, 10, 12, 46, 12, 14, 16], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(sellers, filename: str) -> http.HttpResponse:
    data = build_xlsx(sellers)
    resp = http.HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def export_job_xlsx(job_id: int) -> http.HttpResponse:
    job = CollectionJob.objects.get(pk=job_id)
    sellers = Seller.objects.filter(collection_links__job=job).distinct()
    if job.cities.exists():
        sellers = sellers.filter(city__in=job.cities.all()).distinct()
    return xlsx_response(sellers, f"job-{job_id}-sellers.xlsx")


def export_all_xlsx() -> http.HttpResponse:
    return xlsx_response(Seller.objects.all().prefetch_related("contacts", "categories").select_related("city"), "all-sellers.xlsx")
