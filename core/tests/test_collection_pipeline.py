from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from core.adapters.base import CollectionBlocked, MarketplaceAdapter, SellerData
from core.export import export_job_xlsx
from core.models import Category, City, CollectionJob, CollectionJobSeller, Marketplace, Seller


class FakeStreamingAdapter(MarketplaceAdapter):
    code = Marketplace.OZON

    def __init__(self, mode="normal"):
        self.mode = mode
        self.iter_calls = 0
        self.detail_calls = 0
        self.starts = []
        self.block_once = mode == "block_detail_once"

    def get_categories(self):
        return []

    def iter_sellers(self, category, city=None, max_sellers=0, start_page=1, start_cursor=None):
        self.iter_calls += 1
        self.starts.append((start_page, start_cursor))
        if self.mode == "blocked_discovery":
            yield SellerData(Marketplace.OZON, "before-block", name="Before"), {
                "page": 2, "cursor": "/search?page=2", "finished": False,
            }
            yield None, {"page": 2, "cursor": "/search?page=2", "finished": False}
            raise CollectionBlocked("HTTP 403")
        if self.mode == "empty":
            yield None, {"page": 1, "cursor": None, "finished": True}
            return
        yield SellerData(Marketplace.OZON, "seller-1", name="Discovered"), {
            "page": 2, "cursor": None, "finished": True,
        }
        yield None, {"page": 2, "cursor": None, "finished": True}

    def fetch_seller(self, ref):
        self.detail_calls += 1
        if self.block_once:
            self.block_once = False
            raise CollectionBlocked("detail blocked")
        return SellerData(
            Marketplace.OZON,
            ref,
            name="Detailed",
            inn="7801234567",
            legal_address="г. Уфа, ул. Ленина, 1",
        )


class CollectionPipelineTest(TestCase):
    def make_job(self, cities=None):
        category = Category.objects.create(marketplace=Marketplace.OZON, external_id="q", title="Q")
        job = CollectionJob.objects.create(marketplace=Marketplace.OZON)
        job.categories.set([category])
        if cities:
            job.cities.set(cities)
        return job

    def test_streaming_keeps_previous_page_before_block(self):
        job = self.make_job()
        adapter = FakeStreamingAdapter("blocked_discovery")
        with patch("core.tasks.get_adapter", return_value=adapter):
            from core.tasks import run_collection_job

            run_collection_job(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, CollectionJob.Status.PAUSED)
        self.assertEqual(job.source_status, "blocked")
        self.assertEqual(job.job_sellers.count(), 1)
        self.assertEqual(job.checkpoint["discovery"][f"ozon:{job.categories.first().id}:0"]["cursor"], "/search?page=2")
        adapter.mode = "normal"
        job.status = CollectionJob.Status.QUEUED
        job.save(update_fields=["status"])
        with patch("core.tasks.get_adapter", return_value=adapter):
            from core.tasks import run_collection_job

            run_collection_job(job.id)
        self.assertEqual(adapter.starts[-1], (2, "/search?page=2"))

    def test_blocked_detail_is_retryable_on_resume(self):
        job = self.make_job()
        adapter = FakeStreamingAdapter("block_detail_once")
        with patch("core.tasks.get_adapter", return_value=adapter):
            from core.tasks import run_collection_job

            run_collection_job(job.id)
            link = CollectionJobSeller.objects.get(job=job)
            self.assertEqual(link.detail_status, "blocked")
            job.refresh_from_db()
            self.assertEqual(job.status, CollectionJob.Status.PAUSED)
            job.status = CollectionJob.Status.QUEUED
            job.save(update_fields=["status"])
            run_collection_job(job.id)
        link.refresh_from_db()
        job.refresh_from_db()
        self.assertEqual(link.detail_status, "done")
        self.assertEqual(job.status, CollectionJob.Status.COMPLETED)

    def test_empty_is_completed_empty(self):
        job = self.make_job()
        adapter = FakeStreamingAdapter("empty")
        with patch("core.tasks.get_adapter", return_value=adapter):
            from core.tasks import run_collection_job

            run_collection_job(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, CollectionJob.Status.COMPLETED)
        self.assertEqual(job.source_status, "empty")

    def test_non_wb_multi_city_does_one_discovery(self):
        cities = [City.objects.get_or_create(name=name)[0] for name in ("Уфа", "Челябинск", "Екатеринбург")]
        job = self.make_job(cities)
        adapter = FakeStreamingAdapter()
        with patch("core.tasks.get_adapter", return_value=adapter):
            from core.tasks import run_collection_job

            run_collection_job(job.id)
        self.assertEqual(adapter.iter_calls, 1)

    def test_dadata_runs_after_detail_inn(self):
        job = self.make_job()
        adapter = FakeStreamingAdapter()
        with patch("core.tasks.get_adapter", return_value=adapter), patch("core.tasks.enrich_with_dadata") as enrich:
            from core.tasks import run_collection_job

            run_collection_job(job.id)
        seller = Seller.objects.get(external_seller_id="seller-1")
        enrich.assert_called_once_with(seller)

    def test_job_export_is_limited_to_job_links(self):
        job = self.make_job()
        first = Seller.objects.create(marketplace=Marketplace.OZON, external_seller_id="one", name="One")
        Seller.objects.create(marketplace=Marketplace.OZON, external_seller_id="two", name="Two")
        CollectionJobSeller.objects.create(job=job, seller=first, external_seller_id="one")
        response = export_job_xlsx(job.id)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        values = [row[2] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(values, ["One"])
